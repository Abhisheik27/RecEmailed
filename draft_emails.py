"""
draft_emails.py
---------------
Automates creating Gmail draft emails for recruiter outreach.

How it works (step by step):
    1. Reads recruiter data (Name, Email, Company, Role) from an Excel file.
    2. Fills a predefined email template with each recruiter's data.
    3. Attaches your resume PDF to every draft.
    4. Creates a draft in your Gmail inbox — does NOT send.
    5. You manually review and hit send yourself.

Authentication:
    - Uses Gmail API with OAuth2.
    - On the FIRST run, a browser window opens for Google login.
    - After that, a token.json file is cached so you won't need to log in again.

Prerequisites:
    - credentials.json  → downloaded from Google Cloud Console
    - recruiters.xlsx    → Excel file with columns: Name, Email, Company, Role
    - resume.pdf         → your resume file

Usage:
    python draft_emails.py
"""

import base64
import os
import re
import sys
from email.mime.application import MIMEApplication
from email.mime.multipart import MIMEMultipart
from email.mime.text import MIMEText

from google.auth.transport.requests import Request
from google.oauth2.credentials import Credentials
from google_auth_oauthlib.flow import InstalledAppFlow
from googleapiclient.discovery import build
from openpyxl import load_workbook


# ---------------------------------------------------------------------------
# CONFIGURATION — edit these values to match your setup
# ---------------------------------------------------------------------------

# Path to the Excel file with recruiter data
EXCEL_FILE = "recruiters.xlsx"

# Path to your resume PDF
RESUME_FILE = "resume.pdf"

# Gmail API OAuth credentials file (from Google Cloud Console)
CREDENTIALS_FILE = "credentials.json"

# Cached token file (auto-created after first login)
TOKEN_FILE = "token.json"

# Gmail API scope — "compose" lets us create drafts without full mailbox access
SCOPES = ["https://www.googleapis.com/auth/gmail.compose"]

# Your email subject line template (placeholders get replaced per recruiter)
SUBJECT_TEMPLATE = "Interested in {Role} opportunity at {Company}"

# ---------------------------------------------------------------------------
# EMAIL TEMPLATE — customise this to your liking!
# ---------------------------------------------------------------------------
# Placeholders available:  {Name}  {Company}  {Role}
#
# Example:
#   "Hi {Name},\n\nI'm excited about the {Role} role at {Company}..."
#
# Keep the triple-quoted string so you can write multi-line emails easily.
# ---------------------------------------------------------------------------

EMAIL_TEMPLATE = """\
Hi {Name},

I hope this message finds you well! I came across the {Role} position \
at {Company} and I'm very excited about the opportunity.

I believe my skills and experience align well with what your team is \
looking for. I've attached my resume for your reference and would love \
the chance to discuss how I can contribute to {Company}.

Please let me know if there's a convenient time to chat. I'm happy to \
work around your schedule.

Thank you for your time, and I look forward to hearing from you!

Best regards,
[Your Name]
"""


# ---------------------------------------------------------------------------
# HELPER FUNCTIONS
# ---------------------------------------------------------------------------


def authenticate_gmail():
    """
    Authenticate with the Gmail API using OAuth2.

    How it works:
        - If token.json exists and is still valid → reuse it (no login needed).
        - If the token is expired but has a refresh token → refresh silently.
        - Otherwise → open a browser window for Google login.

    Returns:
        A Gmail API service object you can call methods on.

    Think of it like logging into Gmail, but programmatically.
    """
    creds = None

    # Step 1: Check if we already have a saved token
    if os.path.exists(TOKEN_FILE):
        creds = Credentials.from_authorized_user_file(TOKEN_FILE, SCOPES)

    # Step 2: If no valid credentials, get new ones
    if not creds or not creds.valid:
        if creds and creds.expired and creds.refresh_token:
            # Token expired but we can refresh it silently
            print("🔄 Refreshing expired token...")
            creds.refresh(Request())
        else:
            # No token at all — need full login flow
            if not os.path.exists(CREDENTIALS_FILE):
                print(
                    f"❌ Error: '{CREDENTIALS_FILE}' not found.\n"
                    "   Download it from Google Cloud Console → APIs & Services → Credentials.\n"
                    "   See README.md for step-by-step instructions."
                )
                sys.exit(1)

            print("🌐 Opening browser for Google login...")
            flow = InstalledAppFlow.from_client_secrets_file(CREDENTIALS_FILE, SCOPES)
            creds = flow.run_local_server(port=0)

        # Step 3: Save the token for next time
        with open(TOKEN_FILE, "w") as token_file:
            token_file.write(creds.to_json())
        print("💾 Token saved to token.json (you won't need to login again).")

    # Step 4: Build and return the Gmail service object
    service = build("gmail", "v1", credentials=creds)
    return service


def validate_email(email: str) -> bool:
    """
    Check if an email address looks valid using a simple regex.

    Examples:
        validate_email("jane@google.com")    → True
        validate_email("not-an-email")       → False
        validate_email("")                   → False

    This is NOT a full RFC-compliant check, but catches obvious typos.
    """
    pattern = r"^[a-zA-Z0-9._%+\-]+@[a-zA-Z0-9.\-]+\.[a-zA-Z]{2,}$"
    return bool(re.match(pattern, email))


def read_recruiter_data(filepath: str) -> list[dict]:
    """
    Read recruiter data from an Excel file.

    Expected columns (first row = headers):
        Name | Email | Company | Role

    Returns:
        A list of dictionaries, one per recruiter. Example:
        [
            {"Name": "Jane", "Email": "jane@google.com", "Company": "Google", "Role": "SWE"},
            {"Name": "John", "Email": "john@meta.com",   "Company": "Meta",   "Role": "Backend Dev"},
        ]

    Raises SystemExit if the file is missing or has wrong columns.
    """
    if not os.path.exists(filepath):
        print(f"❌ Error: Excel file '{filepath}' not found.")
        sys.exit(1)

    wb = load_workbook(filepath, read_only=True)
    ws = wb.active

    # Read all rows into a list (first row = headers)
    rows = list(ws.iter_rows(values_only=True))
    wb.close()

    if len(rows) < 2:
        print("❌ Error: Excel file has no data rows (only headers or empty).")
        sys.exit(1)

    # Validate headers
    expected_headers = {"name", "email", "company", "role"}
    actual_headers = {str(h).strip().lower() for h in rows[0] if h is not None}

    if not expected_headers.issubset(actual_headers):
        missing = expected_headers - actual_headers
        print(
            f"❌ Error: Excel file is missing required columns: {missing}\n"
            f"   Expected: Name, Email, Company, Role\n"
            f"   Found:    {[str(h) for h in rows[0]]}"
        )
        sys.exit(1)

    # Build a mapping from lowercase header → column index
    header_map = {}
    for idx, header in enumerate(rows[0]):
        if header is not None:
            header_map[str(header).strip().lower()] = idx

    # Parse each data row into a dictionary
    recruiters = []
    for row_num, row in enumerate(rows[1:], start=2):
        recruiter = {}
        skip = False

        for field in ["name", "email", "company", "role"]:
            col_idx = header_map[field]
            value = row[col_idx] if col_idx < len(row) else None

            if value is None or str(value).strip() == "":
                print(f"⚠️  Row {row_num}: Missing '{field}' — skipping this row.")
                skip = True
                break

            recruiter[field.capitalize()] = str(value).strip()

        if skip:
            continue

        # Validate email format
        if not validate_email(recruiter["Email"]):
            print(f"⚠️  Row {row_num}: Invalid email '{recruiter['Email']}' — skipping.")
            continue

        recruiters.append(recruiter)

    if not recruiters:
        print("❌ Error: No valid recruiter rows found in Excel file.")
        sys.exit(1)

    print(f"📋 Loaded {len(recruiters)} recruiter(s) from '{filepath}'.\n")
    return recruiters


def build_email_message(
    to_email: str,
    subject: str,
    body: str,
    resume_path: str,
) -> dict:
    """
    Build a Gmail-compatible email message with a PDF attachment.

    How it works:
        1. Create a multipart MIME message (like an envelope with multiple things inside).
        2. Add the email body as plain text.
        3. Read the resume PDF and attach it as a binary attachment.
        4. Encode the whole thing in base64 (Gmail API requirement).

    Args:
        to_email:    Recipient's email address.
        subject:     Email subject line.
        body:        Email body text.
        resume_path: Path to the resume PDF file.

    Returns:
        A dict with a "raw" key containing the base64-encoded email.
        This is the format Gmail API expects for creating drafts.
    """
    # Create the multipart message (think of it as an envelope)
    message = MIMEMultipart()
    message["to"] = to_email
    message["subject"] = subject

    # Attach the email body as plain text
    message.attach(MIMEText(body, "plain"))

    # Read and attach the resume PDF
    with open(resume_path, "rb") as pdf_file:
        pdf_attachment = MIMEApplication(pdf_file.read(), _subtype="pdf")
        pdf_attachment.add_header(
            "Content-Disposition",
            "attachment",
            filename=os.path.basename(resume_path),
        )
        message.attach(pdf_attachment)

    # Encode the entire message in URL-safe base64 (Gmail API requirement)
    raw_bytes = message.as_bytes()
    encoded_message = base64.urlsafe_b64encode(raw_bytes).decode("utf-8")

    return {"raw": encoded_message}


# ---------------------------------------------------------------------------
# MAIN — orchestrates the entire workflow
# ---------------------------------------------------------------------------


def main():
    """
    Main entry point. Orchestrates the full workflow:
        1. Validate that the resume file exists.
        2. Read recruiter data from Excel.
        3. Authenticate with Gmail.
        4. For each recruiter, build an email and create a draft.
        5. Print a summary.
    """
    print("=" * 60)
    print("  📧  Gmail Draft Creator — Recruiter Outreach")
    print("=" * 60)
    print()

    # --- Step 1: Validate resume file exists ---
    if not os.path.exists(RESUME_FILE):
        print(
            f"❌ Error: Resume file '{RESUME_FILE}' not found.\n"
            f"   Place your resume PDF in the same directory as this script."
        )
        sys.exit(1)
    print(f"📎 Resume file found: {RESUME_FILE}")

    # --- Step 2: Read recruiter data from Excel ---
    recruiters = read_recruiter_data(EXCEL_FILE)

    # --- Step 3: Authenticate with Gmail ---
    print("🔐 Authenticating with Gmail API...")
    service = authenticate_gmail()
    print("✅ Authentication successful!\n")

    # --- Step 4: Create a draft for each recruiter ---
    success_count = 0
    fail_count = 0

    for i, recruiter in enumerate(recruiters, start=1):
        name = recruiter["Name"]
        email = recruiter["Email"]
        company = recruiter["Company"]
        role = recruiter["Role"]

        print(f"[{i}/{len(recruiters)}] Creating draft for {name} ({email})...")

        try:
            # Fill the template with this recruiter's data
            subject = SUBJECT_TEMPLATE.format(
                Name=name, Company=company, Role=role
            )
            body = EMAIL_TEMPLATE.format(
                Name=name, Company=company, Role=role
            )

            # Build the MIME email with resume attached
            message_body = build_email_message(
                to_email=email,
                subject=subject,
                body=body,
                resume_path=RESUME_FILE,
            )

            # Create the draft in Gmail (synchronous call via Google API client)
            draft = (
                service.users()
                .drafts()
                .create(userId="me", body={"message": message_body})
                .execute()
            )

            draft_id = draft.get("id", "unknown")
            print(f"   ✅ Draft created (ID: {draft_id})")
            success_count += 1

        except Exception as e:
            print(f"   ❌ Failed to create draft: {e}")
            fail_count += 1

    # --- Step 5: Print summary ---
    print()
    print("=" * 60)
    print("  📊  Summary")
    print("=" * 60)
    print(f"  ✅ Drafts created: {success_count}")
    print(f"  ❌ Failed:         {fail_count}")
    print(f"  📧 Total:          {len(recruiters)}")
    print()

    if success_count > 0:
        print("👉 Open Gmail → Drafts to review and send your emails!")
    print()


if __name__ == "__main__":
    main()
